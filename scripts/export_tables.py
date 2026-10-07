"""Export every table in manuscript/FINAL_INTEGRATED_MANUSCRIPT.docx to CSV:
main tables to manuscript/tables/, supplementary tables to supplementary/tables/."""
import csv
import glob
import os
import re
import docx

d = docx.Document("manuscript/FINAL_INTEGRATED_MANUSCRIPT.docx")
caps = [p.text for p in d.paragraphs if re.match(r"^(Table \d+\.|Supplementary Table S\d+\.)", p.text)]
assert len(caps) == len(d.tables), (len(caps), len(d.tables))
for f in glob.glob("manuscript/tables/*.csv") + glob.glob("supplementary/tables/*.csv"):
    os.remove(f)
os.makedirs("manuscript/tables", exist_ok=True); os.makedirs("supplementary/tables", exist_ok=True)
for cap, table in zip(caps, d.tables):
    key = re.match(r"^(Table \d+|Supplementary Table S\d+)", cap).group(1)
    title = re.sub(r"[^A-Za-z0-9]+", "_", cap.split(".", 1)[1].strip().split(".")[0])[:60].strip("_")
    folder = "manuscript/tables" if key.startswith("Table") else "supplementary/tables"
    path = f"{folder}/{key.replace(' ', '_')}_{title}.csv"
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        w = csv.writer(fh)
        for r in table.rows:
            w.writerow([c.text for c in r.cells])
    print(path)
